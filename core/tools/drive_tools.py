from __future__ import annotations

import io
import re

from googleapiclient.discovery import build
from googleapiclient.http import MediaIoBaseDownload, MediaIoBaseUpload
from pypdf import PdfReader

from core.google_auth import get_credentials
from core.tools import Tool, register

_NOT_CONNECTED = {
    "error": "Drive no está conectado. Pedile al usuario que lo conecte en Configuración → Conexiones → Google."
}

_XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

_EXPORT_MIME_MAP = {
    "application/vnd.google-apps.document": "text/plain",
    # xlsx rather than csv: the csv export only carries the FIRST tab
    "application/vnd.google-apps.spreadsheet": _XLSX_MIME,
    "application/vnd.google-apps.presentation": "text/plain",
}

# What a tool result may carry back to a small local model without flooding its context.
MAX_READ_CHARS = 12000


def extract_file_id(value: str) -> str:
    """Accepts a bare id or any Google Docs/Sheets/Slides/Drive link."""
    value = value.strip()
    match = re.search(r"/d/([\w-]{20,})", value) or re.search(r"[?&]id=([\w-]{20,})", value)
    return match.group(1) if match else value


def _spreadsheet_to_text(raw_bytes: bytes) -> str:
    """Every tab as plain rows ('a | b | c'), each under a '## Hoja: name' heading."""
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(raw_bytes), data_only=True, read_only=True)
    parts = []
    for sheet in workbook.worksheets:
        rows = []
        for row in sheet.iter_rows(values_only=True):
            cells = ["" if cell is None else str(cell).replace("\n", " ") for cell in row]
            if any(cells):
                rows.append(" | ".join(cells).rstrip(" |"))
        parts.append(f"## Hoja: {sheet.title}\n" + "\n".join(rows))
    return "\n\n".join(parts)


def search_drive_files(query: str) -> dict:
    credentials = get_credentials()
    if credentials is None:
        return _NOT_CONNECTED
    service = build("drive", "v3", credentials=credentials)
    escaped = query.replace("'", "\\'")
    result = service.files().list(
        q=f"name contains '{escaped}' and trashed = false",
        fields="files(id,name,mimeType,modifiedTime)",
        pageSize=10,
    ).execute()
    return {"results": result.get("files", [])}


def read_drive_file(file_id: str) -> dict:
    file_id = extract_file_id(file_id)
    credentials = get_credentials()
    if credentials is None:
        return _NOT_CONNECTED
    service = build("drive", "v3", credentials=credentials)
    meta = service.files().get(fileId=file_id, fields="id,name,mimeType").execute()
    mime_type = meta["mimeType"]

    export_mime = _EXPORT_MIME_MAP.get(mime_type)
    if export_mime:
        raw = service.files().export(fileId=file_id, mimeType=export_mime).execute()
        if export_mime == _XLSX_MIME:
            content = _spreadsheet_to_text(raw)
        else:
            content = raw.decode("utf-8", errors="replace") if isinstance(raw, bytes) else raw
    else:
        buffer = io.BytesIO()
        downloader = MediaIoBaseDownload(buffer, service.files().get_media(fileId=file_id))
        done = False
        while not done:
            _, done = downloader.next_chunk()
        raw_bytes = buffer.getvalue()
        if mime_type == "application/pdf":
            reader = PdfReader(io.BytesIO(raw_bytes))
            content = "\n\n".join(page.extract_text() or "" for page in reader.pages)
        else:
            content = raw_bytes.decode("utf-8", errors="replace")

    result = {"id": file_id, "name": meta["name"], "mime_type": mime_type, "content": content}
    if len(content) > MAX_READ_CHARS:
        result["content"] = content[:MAX_READ_CHARS]
        result["truncated"] = f"el archivo tiene {len(content)} caracteres; se muestran los primeros {MAX_READ_CHARS}"
    return result


def write_drive_file(name: str, content: str, file_id: str | None = None, folder_id: str | None = None) -> dict:
    credentials = get_credentials()
    if credentials is None:
        return _NOT_CONNECTED
    service = build("drive", "v3", credentials=credentials)
    media = MediaIoBaseUpload(io.BytesIO(content.encode("utf-8")), mimetype="text/plain")

    if file_id:
        updated = service.files().update(fileId=file_id, media_body=media, body={"name": name}).execute()
        return {"id": updated["id"], "name": name, "overwritten": True}

    metadata = {"name": name}
    if folder_id:
        metadata["parents"] = [folder_id]
    created = service.files().create(body=metadata, media_body=media, fields="id,name").execute()
    return {"id": created["id"], "name": created["name"], "created": True}


register(Tool(
    name="search_drive_files",
    description="Search Google Drive files by name.",
    parameters={
        "type": "object",
        "properties": {"query": {"type": "string", "description": "Text to search for in file names"}},
        "required": ["query"],
    },
    func=search_drive_files,
))

register(Tool(
    name="read_drive_file",
    description=(
        "Read the text content of a Google Doc, Sheet (every tab), Slides, PDF or text file. Pass the file id OR the "
        "full link the user pasted (docs.google.com/... or drive.google.com/...): never say you can't open links, "
        "call this tool with it. Long files come back cut to the first part (a 'truncated' note says so)."
    ),
    parameters={
        "type": "object",
        "properties": {"file_id": {"type": "string", "description": "Drive file id (from search_drive_files) or a full Google Docs/Sheets/Drive URL"}},
        "required": ["file_id"],
    },
    func=read_drive_file,
))

register(Tool(
    name="write_drive_file",
    description=(
        "Create a new text file in Drive, or overwrite an existing one if file_id is given. "
        "Requires user confirmation."
    ),
    parameters={
        "type": "object",
        "properties": {
            "name": {"type": "string", "description": "File name"},
            "content": {"type": "string", "description": "Full text content to write"},
            "file_id": {"type": "string", "description": "Existing file id to overwrite, omit to create a new file"},
            "folder_id": {"type": "string", "description": "Parent folder id for new files, omit for root"},
        },
        "required": ["name", "content"],
    },
    func=write_drive_file,
    requires_confirmation=True,
))
